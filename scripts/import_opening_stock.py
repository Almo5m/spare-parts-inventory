#!/usr/bin/env python3
"""
استيراد المخزون الافتتاحي من شيتات Excel (مخزنين) إلى Supabase.

السكريبت مبيكتبش في القاعدة. بيقرا الشيتين، وبيطلّع في فولدر out/:
  - opening-stock.sql  : ملف SQL يتنفّذ مرة واحدة من SQL Editor (أمر DO واحد، كله ذرّي)
  - import-review.xlsx : ملف مراجعة (المنتجات والكميات لكل مخزن + اللي اتخطّاه + فروق الأسعار)

شكل كل شيت: عمود الاسم، عمود الكمية بالقطعة، عمود السعر (سعر البيع). أول صف عناوين.
أي صف من غير كمية وسعر (زي عناوين الأقسام) بيتخطّى ويتسجّل في تقرير المراجعة.

الاستخدام:
    pip install openpyxl
    python scripts/import_opening_stock.py --main "مخزن1.xlsx" --branch "مخزن2.xlsx"
    python scripts/import_opening_stock.py --main ... --branch ... --merge merge.json --prices prices.json

merge.json (اختياري): {"اسم الصنف في المخزن الفرعي": "نفس الصنف باسمه في المخزن الرئيسي"}
للأصناف اللي أسماءها مختلفة بس هي نفس المنتج. من غيره، أي اسمين مختلفين بيتعاملوا كمنتجين.

prices.json (اختياري): {"اسم الصنف": 475}
بيوحّد سعر الصنف في كل المخازن (لما الشيتين فيهم سعرين مختلفين لنفس الصنف وانت قررت أنهي واحد).

تنبيه خصوصية: فولدر out/ فيه بيانات العميل (أصناف وأسعار وكميات). ما ترفعوش على GitHub
(اتضاف في .gitignore).
"""
import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

WAREHOUSE_NAMES = ("مخزن 1", "مخزن 2")  # الأول هو الافتراضي (الأساسي)، وتتغيّر أسماؤهم من الواجهة بعدين
OPENING_NOTE = "رصيد افتتاحي — استيراد من الشيتات"
CODE_PREFIX = "P-"

_AR_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")


def clean(text) -> str:
    """تنضيف الاسم: شيل التشكيل والتطويل وكرر المسافات. مبنغيّرش الإملاء."""
    s = str(text).translate(_AR_DIGITS)
    s = re.sub(r"[\u064B-\u0652\u0640]", "", s)
    return re.sub(r"\s+", " ", s).strip()


def match_key(text) -> str:
    """مفتاح المطابقة بين الشيتين: بيوحّد الهمزات والياء والتاء المربوطة بس للمقارنة."""
    s = clean(text).lower()
    s = re.sub("[أإآ]", "ا", s).replace("ى", "ي").replace("ة", "ه")
    return s


def to_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).translate(_AR_DIGITS).replace(",", "").strip()
    try:
        return float(text)
    except ValueError:
        return None


def read_sheet(path: str):
    """بيرجّع (الأصناف، الصفوف المتخطّاة). الصف الصالح لازم يكون فيه اسم وكمية ≥ 0 وسعر ≥ 0."""
    ws = load_workbook(path, data_only=True).active
    items, skipped = [], []
    for row_no, (name, qty, price) in enumerate(ws.iter_rows(min_row=2, max_col=3, values_only=True), start=2):
        if name is None and qty is None and price is None:
            continue
        q, p = to_number(qty), to_number(price)
        if not name or q is None or p is None or q < 0 or p < 0:
            skipped.append({"row": row_no, "name": clean(name) if name else "", "qty": qty, "price": price})
            continue
        items.append({"row": row_no, "name": clean(name), "qty": q, "price": p})
    return items, skipped


def fmt(n: float) -> str:
    return str(int(n)) if float(n).is_integer() else repr(round(n, 4))


def sql_text(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def build(main_items, branch_items, merge_map):
    products = {}  # key -> dict
    order = []

    def add(item, wh):
        key = match_key(item["name"])
        if key not in products:
            products[key] = {"name": item["name"], "main": None, "branch": None}
            order.append(key)
        slot = products[key]
        if slot[wh] is not None:
            raise SystemExit(f"الاسم ده مكرر في نفس الشيت ({wh}): {item['name']} (صف {item['row']})")
        slot[wh] = item

    for it in main_items:
        add(it, "main")

    main_keys = {match_key(it["name"]) for it in main_items}
    merge_keys = {match_key(k): match_key(v) for k, v in merge_map.items()}
    for target in merge_keys.values():
        if target not in main_keys:
            raise SystemExit(f"في merge.json اسم مش موجود في المخزن الرئيسي: {target}")

    for it in branch_items:
        k = match_key(it["name"])
        if k in merge_keys:
            it = {**it, "name": products[merge_keys[k]]["name"]}
        add(it, "branch")

    rows = []
    for i, key in enumerate(order, start=1):
        p = products[key]
        rows.append({
            "code": f"{CODE_PREFIX}{i:05d}",
            "name": p["name"],
            "main": p["main"],
            "branch": p["branch"],
        })
    return rows


def apply_price_overrides(rows, overrides):
    """بيوحّد سعر الأصناف المحددة في كل المخازن اللي فيها الصنف."""
    by_key = {match_key(r["name"]): r for r in rows}
    for name, price in overrides.items():
        row = by_key.get(match_key(name))
        if row is None:
            raise SystemExit(f"في prices.json اسم مش موجود في الشيتات: {name}")
        if not isinstance(price, (int, float)) or price < 0:
            raise SystemExit(f"السعر لازم يكون رقم صفر أو أكبر: {name} = {price}")
        for slot in ("main", "branch"):
            if row[slot]:
                row[slot] = {**row[slot], "price": float(price)}


def write_sql(rows, out_path: Path):
    """
    الملف كله أمر DO واحد (مفيش begin/commit ولا جداول مؤقتة). السبب: الـ SQL Editor في Supabase
    مش بيضمن إن الأوامر المتتالية تتنفذ في نفس الجلسة والـ transaction، فالجدول المؤقت كان
    بيتمسح قبل ما الأمر اللي بعده يستخدمه. أمر DO واحد بيتنفذ ذرّيًا: يا كله يتسجّل يا ولا حاجة.
    """
    main_rows = [r for r in rows if r["main"]]
    branch_rows = [r for r in rows if r["branch"]]
    main_qty = sum(r["main"]["qty"] for r in main_rows)
    branch_qty = sum(r["branch"]["qty"] for r in branch_rows)

    data = [
        {
            "code": r["code"],
            "name": r["name"],
            "p1": r["main"]["price"] if r["main"] else None,
            "q1": r["main"]["qty"] if r["main"] else None,
            "p2": r["branch"]["price"] if r["branch"] else None,
            "q2": r["branch"]["qty"] if r["branch"] else None,
        }
        for r in rows
    ]
    # سطر لكل منتج عشان الملف يتقرا ويتراجع بسهولة
    json_lines = ",\n".join("  " + json.dumps(d, ensure_ascii=False) for d in data)
    if "$json$" in json_lines:
        raise SystemExit("اسم صنف فيه الرمز $json$ ومينفعش يتحط في الملف")

    recordset = (
        "jsonb_to_recordset(v_data) as r(code text, name text, "
        "p1 numeric, q1 numeric, p2 numeric, q2 numeric)"
    )

    def warehouse_block(n: int) -> str:
        return f"""
  -- {WAREHOUSE_NAMES[n - 1]}: الرصيد + جرد افتتاحي + دفعات + حركات
  insert into public.inventory (product_id, warehouse_id, quantity)
  select p.id, v_wh{n}, r.q{n}
  from {recordset}
  join public.products p on p.part_number = r.code
  where r.q{n} is not null;

  insert into public.stocktake_items (stocktake_id, product_id, system_quantity, counted_quantity, difference)
  select v_st{n}, p.id, 0, r.q{n}, r.q{n}
  from {recordset}
  join public.products p on p.part_number = r.code
  where r.q{n} is not null;

  insert into public.inventory_lots
    (product_id, warehouse_id, source_type, source_id, unit_cost, quantity_received, quantity_remaining, default_sale_price)
  select p.id, v_wh{n}, 'stocktake', v_st{n}, 0, r.q{n}, r.q{n}, r.p{n}
  from {recordset}
  join public.products p on p.part_number = r.code
  where r.q{n} is not null;

  insert into public.inventory_movements (product_id, warehouse_id, movement_type, quantity, reference_id)
  select p.id, v_wh{n}, 'adjustment'::public.inventory_movement_type, r.q{n}, v_st{n}
  from {recordset}
  join public.products p on p.part_number = r.code
  where r.q{n} is not null;
"""

    sql = f"""-- ============================================================================
-- استيراد المخزون الافتتاحي (اتولّد بـ scripts/import_opening_stock.py)
-- {len(rows)} منتج | {WAREHOUSE_NAMES[0]}: {len(main_rows)} صنف ({fmt(main_qty)} قطعة) | {WAREHOUSE_NAMES[1]}: {len(branch_rows)} صنف ({fmt(branch_qty)} قطعة)
--
-- شغّله مرة واحدة من Supabase > SQL Editor، على قاعدة فاضية (بعد تطبيق كل الـ migrations).
-- الملف كله أمر DO واحد: لو أي حاجة اتخالفت (أعداد أو مجاميع) كله بيتلغي ومفيش حاجة بتتسجّل.
-- ماتضيفش begin/commit ومتقسّمهوش على أكتر من تشغيل.
-- سعر الشراء = صفر لكل الأصناف (مش معروف لسه). الدفعات الافتتاحية معلّمة بالملاحظة:
--   '{OPENING_NOTE}'
-- عشان نصلّح تكلفتها بعدين بسكريبت إعادة التكلفة.
-- سعر البيع على المنتج = سعر {WAREHOUSE_NAMES[0]} (أو {WAREHOUSE_NAMES[1]} لو الصنف فيه بس).
-- سعر كل مخزن محفوظ على دفعاته (default_sale_price) وشاشة البيع بتقترحه.
-- ============================================================================

do $import$
declare
  v_wh1 uuid; v_wh2 uuid;
  v_st1 uuid; v_st2 uuid;
  v_products int;
  v_inv_1 int; v_inv_2 int;
  v_qty_1 numeric; v_qty_2 numeric;
  v_lots int; v_mov int;
  -- كل منتج: كود، اسم، (سعر، كمية) مخزن 1، (سعر، كمية) مخزن 2
  v_data jsonb := $json$[
{json_lines}
]$json$::jsonb;
begin
  if exists (select 1 from public.products) or exists (select 1 from public.warehouses) then
    raise exception 'القاعدة فيها منتجات أو مخازن بالفعل. الملف ده للاستيراد الأول بس';
  end if;

  -- المخازن
  insert into public.warehouses (name, is_default) values ({sql_text(WAREHOUSE_NAMES[0])}, true) returning id into v_wh1;
  insert into public.warehouses (name, is_default) values ({sql_text(WAREHOUSE_NAMES[1])}, false) returning id into v_wh2;

  -- المنتجات
  insert into public.products (part_number, name, base_unit, cost_price, sale_price)
  select r.code, r.name, 'قطعة', 0, coalesce(r.p1, r.p2)
  from {recordset};

  -- جرد افتتاحي لكل مخزن (بيعلّم الدفعات عشان تتصلّح تكلفتها بعدين)
  insert into public.stocktakes (warehouse_id, notes) values (v_wh1, {sql_text(OPENING_NOTE)}) returning id into v_st1;
  insert into public.stocktakes (warehouse_id, notes) values (v_wh2, {sql_text(OPENING_NOTE)}) returning id into v_st2;
{warehouse_block(1)}{warehouse_block(2)}
  -- مطابقة نهائية: لو أي رقم مختلف عن الشيت، العملية كلها بتتلغي
  select count(*) into v_products from public.products;
  select count(*), coalesce(sum(quantity), 0) into v_inv_1, v_qty_1 from public.inventory where warehouse_id = v_wh1;
  select count(*), coalesce(sum(quantity), 0) into v_inv_2, v_qty_2 from public.inventory where warehouse_id = v_wh2;
  select count(*) into v_lots from public.inventory_lots where source_type = 'stocktake';
  select count(*) into v_mov from public.inventory_movements where movement_type = 'adjustment';

  if v_products <> {len(rows)}
     or v_inv_1 <> {len(main_rows)} or v_qty_1 <> {fmt(main_qty)}
     or v_inv_2 <> {len(branch_rows)} or v_qty_2 <> {fmt(branch_qty)}
     or v_lots <> {len(main_rows) + len(branch_rows)} or v_mov <> {len(main_rows) + len(branch_rows)} then
    raise exception 'المطابقة فشلت: منتجات %, مخزن1 % صنف / % قطعة, مخزن2 % صنف / % قطعة',
      v_products, v_inv_1, v_qty_1, v_inv_2, v_qty_2;
  end if;

  raise notice 'تم الاستيراد: % منتج | مخزن 1: % صنف (% قطعة) | مخزن 2: % صنف (% قطعة)',
    v_products, v_inv_1, v_qty_1, v_inv_2, v_qty_2;
end
$import$;

-- بعد التنفيذ: راجع الأعداد من واجهة النظام (صفحة المخازن والمنتجات).
"""
    out_path.write_text(sql, encoding="utf-8")


def write_review(rows, main_skipped, branch_skipped, similar, out_path: Path):
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3A5F")

    def sheet(title, headers, data, widths):
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        for d in data:
            ws.append(d)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)
    sheet("المنتجات", ["كود الصنف", "اسم المنتج", "الوحدة", "سعر الشراء", "سعر البيع", "في مخزن 1", "في مخزن 2"],
          [[r["code"], r["name"], "قطعة", 0,
            (r["main"] or r["branch"])["price"],
            "نعم" if r["main"] else "", "نعم" if r["branch"] else ""] for r in rows],
          [12, 44, 10, 12, 12, 11, 11])
    sheet("مخزن_1", ["كود الصنف", "اسم المنتج", "الكمية بالقطعة", "سعر البيع"],
          [[r["code"], r["name"], r["main"]["qty"], r["main"]["price"]] for r in rows if r["main"]],
          [12, 44, 14, 12])
    sheet("مخزن_2", ["كود الصنف", "اسم المنتج", "الكمية بالقطعة", "سعر البيع"],
          [[r["code"], r["name"], r["branch"]["qty"], r["branch"]["price"]] for r in rows if r["branch"]],
          [12, 44, 14, 12])
    both = [r for r in rows if r["main"] and r["branch"]]
    sheet("في_المخزنين", ["كود الصنف", "اسم المنتج", "كمية مخزن 1", "سعر مخزن 1", "كمية مخزن 2", "سعر مخزن 2", "السعر"],
          [[r["code"], r["name"], r["main"]["qty"], r["main"]["price"], r["branch"]["qty"], r["branch"]["price"],
            "" if r["main"]["price"] == r["branch"]["price"] else "مختلف"] for r in both],
          [12, 44, 12, 12, 12, 12, 10])
    sheet("صفوف_اتخطّت", ["الشيت", "رقم الصف", "المحتوى", "الكمية", "السعر"],
          [["مخزن 1", s["row"], s["name"], s["qty"], s["price"]] for s in main_skipped] +
          [["مخزن 2", s["row"], s["name"], s["qty"], s["price"]] for s in branch_skipped],
          [10, 10, 36, 12, 12])
    sheet("اسماء_متشابهة", ["اسم في مخزن 2", "أقرب اسم في مخزن 1"], similar, [44, 44])
    wb.save(out_path)


def find_similar(main_items, branch_items, merged_names):
    import difflib
    main_keys = {match_key(i["name"]): i["name"] for i in main_items}
    out = []
    for b in branch_items:
        k = match_key(b["name"])
        if k in main_keys or k in merged_names:
            continue
        close = difflib.get_close_matches(k, list(main_keys), n=2, cutoff=0.86)
        if close:
            out.append([b["name"], " | ".join(main_keys[c] for c in close)])
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--main", required=True, help="شيت المخزن الأساسي (مخزن 1)")
    ap.add_argument("--branch", required=True, help="شيت المخزن الفرعي (مخزن 2)")
    ap.add_argument("--merge", help="ملف JSON لدمج أسماء مختلفة لنفس الصنف")
    ap.add_argument("--prices", help="ملف JSON لتوحيد سعر أصناف محددة في كل المخازن")
    ap.add_argument("--out", default="out", help="فولدر الإخراج (الافتراضي out)")
    args = ap.parse_args()

    merge_map = json.loads(Path(args.merge).read_text(encoding="utf-8")) if args.merge else {}

    main_items, main_skipped = read_sheet(args.main)
    branch_items, branch_skipped = read_sheet(args.branch)
    rows = build(main_items, branch_items, merge_map)
    if args.prices:
        apply_price_overrides(rows, json.loads(Path(args.prices).read_text(encoding="utf-8")))
    similar = find_similar(main_items, branch_items, {match_key(k) for k in merge_map})

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    write_sql(rows, out / "opening-stock.sql")
    write_review(rows, main_skipped, branch_skipped, similar, out / "import-review.xlsx")

    both = sum(1 for r in rows if r["main"] and r["branch"])
    print(f"منتجات: {len(rows)} (مخزن 1: {len(main_items)} | مخزن 2: {len(branch_items)} | في الاتنين: {both})")
    print(f"صفوف اتخطّت: مخزن 1 = {len(main_skipped)} | مخزن 2 = {len(branch_skipped)}")
    print(f"أسماء متشابهة محتاجة قرار: {len(similar)}")
    print(f"الملفات في: {out.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
